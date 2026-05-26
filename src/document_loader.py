"""Document loading utilities for uploaded and local PDF and DOCX sources."""

from __future__ import annotations

from functools import lru_cache
from io import BytesIO
import logging
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader

try:
    import pypdfium2 as pdfium
except ImportError:  # pragma: no cover - optional dependency
    pdfium = None

try:
    from rapidocr_onnxruntime import RapidOCR
except ImportError:  # pragma: no cover - optional dependency
    RapidOCR = None


logger = logging.getLogger(__name__)
request_logger = logging.getLogger("uvicorn.error")


PageRecord = dict[str, Any]
_OCR_MIN_TEXT_LENGTH = 24
_OCR_RENDER_SCALE = 2.0


def extract_text_from_pdf_file(file: Any, source_type: str = "uploaded") -> list[PageRecord]:
    """Extract page text from a Streamlit-uploaded PDF-like object."""
    source_name = getattr(file, "name", "uploaded_document.pdf")

    try:
        if hasattr(file, "seek"):
            file.seek(0)
        pdf_bytes = file.read()
    except Exception as exc:
        logger.warning("Failed to read uploaded PDF '%s': %s", source_name, exc)
        return []

    return _extract_pages_from_pdf_bytes(pdf_bytes, source_name=source_name, source_type=source_type)


def extract_text_from_pdf_path(pdf_path: str | Path, source_type: str = "local") -> list[PageRecord]:
    """Extract page text from a PDF file stored on disk."""
    path = Path(pdf_path)

    try:
        pdf_bytes = path.read_bytes()
        return _extract_pages_from_pdf_bytes(
            pdf_bytes,
            source_name=path.name,
            source_type=source_type,
        )
    except Exception as exc:
        logger.warning("Failed to read local PDF '%s': %s", path, exc)
        return []


def extract_text_from_docx_file(file: Any, source_type: str = "uploaded") -> list[PageRecord]:
    """Extract document text from a DOCX-like file object."""
    source_name = getattr(file, "name", "uploaded_document.docx")

    try:
        if hasattr(file, "seek"):
            file.seek(0)
        document = Document(file)
    except Exception as exc:
        logger.warning("Failed to read uploaded DOCX '%s': %s", source_name, exc)
        return []

    return _extract_pages_from_docx_document(
        document,
        source_name=source_name,
        source_type=source_type,
    )


def extract_text_from_docx_path(docx_path: str | Path, source_type: str = "local") -> list[PageRecord]:
    """Extract document text from a DOCX file stored on disk."""
    path = Path(docx_path)

    try:
        document = Document(path)
        return _extract_pages_from_docx_document(
            document,
            source_name=path.name,
            source_type=source_type,
        )
    except Exception as exc:
        logger.warning("Failed to read local DOCX '%s': %s", path, exc)
        return []


def load_supported_documents_from_data_folder(data_dir: str = "data") -> list[PageRecord]:
    """Load and extract text from all supported documents inside the local data folder."""
    data_path = Path(data_dir)
    if not data_path.exists():
        logger.warning("Data directory does not exist: %s", data_path)
        return []

    document_paths = sorted(
        path
        for path in data_path.iterdir()
        if path.is_file() and path.suffix.lower() in {".pdf", ".docx"}
    )
    if not document_paths:
        logger.warning("No supported documents found in data directory: %s", data_path)
        return []

    extracted_pages: list[PageRecord] = []
    for document_path in document_paths:
        suffix = document_path.suffix.lower()
        if suffix == ".pdf":
            extracted_pages.extend(extract_text_from_pdf_path(document_path, source_type="local"))
        elif suffix == ".docx":
            extracted_pages.extend(extract_text_from_docx_path(document_path, source_type="local"))

    return extracted_pages


def summarize_extraction(pages: list[PageRecord]) -> dict[str, Any]:
    """Summarize extracted page records for backend tracking and UI display."""
    source_filenames = sorted({page["source"] for page in pages})
    empty_pages = sum(1 for page in pages if not page.get("text", "").strip())
    ocr_pages = sum(1 for page in pages if page.get("extraction_method") == "ocr")

    pages_per_source: dict[str, int] = {}
    for page in pages:
        source = page["source"]
        pages_per_source[source] = pages_per_source.get(source, 0) + 1

    source_types = sorted({page["source_type"] for page in pages}) if pages else []

    return {
        "document_count": len(source_filenames),
        "total_pages_extracted": len(pages),
        "empty_or_skipped_pages": empty_pages,
        "ocr_pages": ocr_pages,
        "source_filenames": source_filenames,
        "pages_per_source": pages_per_source,
        "source_types": source_types,
    }


def _extract_pages_from_pdf_bytes(
    pdf_bytes: bytes,
    *,
    source_name: str,
    source_type: str,
) -> list[PageRecord]:
    """Extract PDF pages and use OCR only for pages with little or no embedded text."""
    try:
        reader = PdfReader(BytesIO(pdf_bytes))
    except Exception as exc:
        logger.warning("Failed to parse PDF '%s': %s", source_name, exc)
        return []

    ocr_document = _load_pdfium_document(pdf_bytes, source_name=source_name)
    return _extract_pages_from_reader(
        reader,
        source_name=source_name,
        source_type=source_type,
        ocr_document=ocr_document,
    )


def _extract_pages_from_reader(
    reader: PdfReader,
    *,
    source_name: str,
    source_type: str,
    ocr_document: Any | None = None,
) -> list[PageRecord]:
    """Normalize page extraction and keep empty pages as explicit records."""
    extracted_pages: list[PageRecord] = []

    for page_index, page in enumerate(reader.pages, start=1):
        try:
            extracted_text = page.extract_text() or ""
        except Exception as exc:
            logger.warning(
                "Failed to extract page %s from '%s': %s",
                page_index,
                source_name,
                exc,
            )
            extracted_text = ""

        extraction_method = "embedded_text"
        normalized_text = extracted_text.strip()
        if _should_use_ocr(normalized_text):
            ocr_text = _extract_text_with_ocr(
                ocr_document,
                page_index=page_index,
                source_name=source_name,
            )
            if ocr_text:
                normalized_text = ocr_text
                extraction_method = "ocr"

        request_logger.info(
            "Extraction method for '%s' page %s: %s",
            source_name,
            page_index,
            extraction_method,
        )

        extracted_pages.append(
            {
                "source": source_name,
                "page_number": page_index,
                "text": normalized_text,
                "source_type": source_type,
                "extraction_method": extraction_method,
            }
        )

    if not extracted_pages:
        logger.warning("PDF '%s' contains no readable pages.", source_name)

    return extracted_pages


def _extract_pages_from_docx_document(
    document: Document,
    *,
    source_name: str,
    source_type: str,
) -> list[PageRecord]:
    """Normalize DOCX extraction into a single page-like record."""
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    extracted_text = "\n\n".join(paragraphs)

    page_record = {
        "source": source_name,
        "page_number": 1,
        "text": extracted_text,
        "source_type": source_type,
    }

    if not extracted_text:
        logger.warning("DOCX '%s' contains no readable paragraphs.", source_name)

    return [page_record]


def _should_use_ocr(text: str) -> bool:
    """Trigger OCR only for pages with effectively no embedded text."""
    compact = " ".join(text.split())
    return len(compact) < _OCR_MIN_TEXT_LENGTH


def _load_pdfium_document(pdf_bytes: bytes, *, source_name: str) -> Any | None:
    """Load a PDFium document for OCR rendering when the optional dependency is available."""
    if pdfium is None:
        return None
    try:
        return pdfium.PdfDocument(pdf_bytes)
    except Exception as exc:
        logger.warning("Failed to prepare OCR renderer for '%s': %s", source_name, exc)
        return None


@lru_cache(maxsize=1)
def _get_ocr_engine() -> Any | None:
    """Build the OCR engine once when the optional dependency is available."""
    if RapidOCR is None:
        return None
    try:
        return RapidOCR()
    except Exception as exc:
        logger.warning("Failed to initialize OCR engine: %s", exc)
        return None


def _extract_text_with_ocr(ocr_document: Any | None, *, page_index: int, source_name: str) -> str:
    """Render a PDF page and run OCR as a fallback when embedded text is missing."""
    if ocr_document is None:
        return ""

    ocr_engine = _get_ocr_engine()
    if ocr_engine is None:
        return ""

    try:
        page = ocr_document[page_index - 1]
        bitmap = page.render(scale=_OCR_RENDER_SCALE)
        image = bitmap.to_pil()
        result, _ = ocr_engine(image)
        lines = [item[1] for item in result or [] if item and len(item) > 1 and item[1]]
        return "\n".join(lines).strip()
    except Exception as exc:
        logger.warning(
            "OCR failed for page %s in '%s': %s",
            page_index,
            source_name,
            exc,
        )
        return ""
